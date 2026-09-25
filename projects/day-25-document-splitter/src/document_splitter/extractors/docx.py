"""DOCX paragraph extraction backed by python-docx."""

from __future__ import annotations

from pathlib import Path
from zipfile import BadZipFile

from document_splitter.extractors.base import ExtractedDocument, DocumentExtractionError


class DocxExtractor:
    """Extract non-empty paragraphs from a DOCX document in order."""

    def extract(self, source: Path) -> ExtractedDocument:
        """Return paragraph text in document order without handling advanced elements."""
        try:
            from docx import Document
            from docx.opc.exceptions import PackageNotFoundError
        except ImportError as error:
            raise DocumentExtractionError(
                "La extracción DOCX requiere instalar python-docx."
            ) from error

        try:
            document = Document(source)
        except (BadZipFile, PackageNotFoundError, OSError, ValueError) as error:
            raise DocumentExtractionError(
                f"No se pudo leer el archivo DOCX: {source}"
            ) from error

        text = "\n".join(
            paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
        )
        if not text.strip():
            raise DocumentExtractionError(
                f"El DOCX no contiene párrafos con texto extraíble: {source}"
            )

        return ExtractedDocument(source=source, text=text)
